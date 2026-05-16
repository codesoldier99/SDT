from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "lectures" / "2026-spring" / "pptx"
NOTES_DIR = ROOT / "lectures" / "2026-spring" / "notes"
ASSET_DIR = ROOT / "lectures" / "2026-spring" / "assets"

COURSE_TITLE = "传感器与检测技术"
COURSE_SUBTITLE = "2026 Spring | 具身感知工程版"

COLORS = {
    "teal": RGBColor(24, 85, 96),
    "teal_dark": RGBColor(16, 59, 67),
    "gold": RGBColor(240, 169, 79),
    "sand": RGBColor(252, 250, 246),
    "mist": RGBColor(239, 244, 244),
    "ink": RGBColor(35, 46, 51),
    "muted": RGBColor(86, 98, 104),
    "line": RGBColor(205, 215, 218),
    "white": RGBColor(255, 255, 255),
}


LECTURES = [
    {
        "id": 1,
        "slug": "lecture-01-introduction",
        "title": "具身感知工程导论",
        "module_type": "讲授",
        "theme": "课程定位与产业场景",
        "objectives": [
            "建立具身感知工程的课程地图，理解传感器在机器人与自动化系统中的角色。",
            "明确精度、噪声、稳定性等工业指标与课程后续实验的对应关系。",
            "理解本课程的数字化资产沉淀要求、GitHub 工作流与综合设计目标。",
        ],
        "foundations": [
            "具身智能系统的感知链路：传感器、前端、估计、控制。",
            "产业场景：无人机、移动机器人、工业检测与定位。",
            "课程评价结构：实验、专报、论文研读、综合设计与资产沉淀。",
        ],
        "principles": [
            "从任务出发定义可观测量。",
            "从指标出发分解传感器选型。",
            "从数据链路出发设计验证闭环。",
            "从版本管理出发沉淀课程资产。",
        ],
        "metrics": [
            ("准确度", "与真值偏差", "决定系统上限"),
            ("重复性", "同条件离散度", "影响稳定输出"),
            ("时延", "采样到决策延迟", "影响实时控制"),
            ("鲁棒性", "环境变化下性能", "决定工程可用性"),
        ],
        "case_title": "案例：移动机器人感知栈",
        "case_points": [
            "GNSS 提供绝对位置基准，IMU 提供高频运动约束。",
            "磁力计和气压计在特定环境下提供航向与高度补偿。",
            "系统性能由硬件、标定、融合算法共同决定。",
        ],
        "practice": [
            "完成开发环境与 GitHub 账号准备。",
            "建立课程仓库、阅读授课计划并确认实验清单。",
            "梳理未来综合设计可能的产品方向。",
        ],
        "homework": [
            "提交课程目标与个人项目兴趣说明。",
            "列出一个感知系统的关键传感器及其核心指标。",
        ],
    },
    {
        "id": 2,
        "slug": "lecture-02-characteristics-datasheet",
        "title": "传感器特性与 Datasheet 分析",
        "module_type": "讲授",
        "theme": "指标、噪声与手册解读",
        "objectives": [
            "掌握静态特性、动态特性与误差来源的分类方法。",
            "能够从 Datasheet 中定位关键参数并判断是否满足任务要求。",
            "理解 Allan 方差与 PSD 在工业 IMU 评估中的作用。",
        ],
        "foundations": [
            "零偏、灵敏度、分辨率、迟滞、线性度。",
            "动态带宽、时延、噪声密度与采样率。",
            "Allan Variance 与随机过程的对应关系。",
        ],
        "principles": [
            "先确定任务精度目标。",
            "再筛选手册中的关键指标。",
            "通过噪声模型估算长期漂移。",
            "结合环境条件做冗余设计。",
        ],
        "metrics": [
            ("噪声密度", "短时抖动", "决定滤波设计"),
            ("Bias Instability", "长期漂移", "决定积分误差"),
            ("Bandwidth", "动态响应", "影响运动捕获"),
            ("Temperature Range", "环境适应性", "影响补偿策略"),
        ],
        "case_title": "案例：三款 IMU 手册对比",
        "case_points": [
            "消费级 IMU 在噪声和温漂上通常弱于工业级器件。",
            "高量程不等于高精度，带宽和噪声需同时考量。",
            "手册中的典型值与最大值都需要关注。",
        ],
        "practice": [
            "完成三款 IMU Datasheet 对比表。",
            "推导 Allan 方差与随机游走项的联系。",
            "为后续标定实验建立指标核对清单。",
        ],
        "homework": [
            "提交 IMU 指标对比报告。",
            "解释一个参数为何会影响导航精度。",
        ],
    },
    {
        "id": 3,
        "slug": "lecture-03-driver-ai-workflow",
        "title": "驱动开发与 AI 协作范式",
        "module_type": "讲授",
        "theme": "嵌入式驱动链路与协作式开发",
        "objectives": [
            "掌握 ESP32 传感器驱动的基本结构与寄存器级调试流程。",
            "理解 AI 工具在代码骨架生成、文档整理和排错中的边界。",
            "建立驱动、实验报告和版本管理之间的统一协作流程。",
        ],
        "foundations": [
            "总线接口：I2C、SPI、UART 的适用场景。",
            "驱动分层：寄存器定义、数据读取、标定补偿与应用接口。",
            "AI 协作工具在嵌入式开发中的高价值环节。",
        ],
        "principles": [
            "阅读手册与时序图。",
            "搭建最小驱动骨架。",
            "通过逻辑分析或串口日志验证。",
            "再用 AI 辅助补全文档与测试。",
        ],
        "metrics": [
            ("初始化成功率", "上电与识别是否稳定", "决定可调试性"),
            ("总线吞吐", "采样能力上限", "影响实时性"),
            ("代码可维护性", "模块边界是否清晰", "影响后续扩展"),
            ("日志可读性", "定位问题效率", "影响实验节奏"),
        ],
        "case_title": "案例：用 AI 生成寄存器驱动骨架",
        "case_points": [
            "AI 可以快速生成读写框架，但寄存器地址和时序必须人工核对。",
            "调试日志模板和错误码设计适合交给 AI 加速。",
            "关键验证仍是硬件链路和示波器结果。",
        ],
        "practice": [
            "建立 ESP32 工程模板与日志输出模块。",
            "尝试用 AI 生成寄存器读写框架并人工审查。",
            "提交一版最小驱动仓库结构。",
        ],
        "homework": [
            "整理一次 AI 协作开发的得失。",
            "补全一个传感器驱动的错误处理流程图。",
        ],
    },
    {
        "id": 4,
        "slug": "lecture-04-hardware-link-lab",
        "title": "实验一：硬件链路打通",
        "module_type": "实验",
        "theme": "从上电到数据输出的闭环验证",
        "objectives": [
            "打通电源、总线、驱动、串口显示的最小实验链路。",
            "掌握常见硬件接线、地址识别与初始化失败的排查方法。",
            "形成实验报告、代码与截图同步归档的习惯。",
        ],
        "foundations": [
            "供电与电平兼容。",
            "I2C 地址扫描与总线错误分析。",
            "串口调试输出的结构化设计。",
        ],
        "principles": [
            "确认供电和引脚映射。",
            "扫描设备地址并识别芯片。",
            "读取原始寄存器并转换数据。",
            "输出稳定日志并保存证据。",
        ],
        "metrics": [
            ("识别成功率", "芯片是否被稳定发现", "反映接线质量"),
            ("首次出数时间", "调通效率", "反映实验准备度"),
            ("串口稳定性", "输出是否连续", "反映程序健壮性"),
            ("记录完整性", "代码与截图是否归档", "反映工程习惯"),
        ],
        "case_title": "案例：I2C 无法识别的排障树",
        "case_points": [
            "优先检查供电与地线，再检查上拉与地址位配置。",
            "利用地址扫描和逻辑分析缩小问题范围。",
            "将排障过程本身写入实验资产。",
        ],
        "practice": [
            "完成地址扫描、芯片识别与原始数据读取。",
            "记录一次初始化失败并复盘排障步骤。",
            "上传实验代码、照片和串口输出截图。",
        ],
        "homework": [
            "提交实验报告。",
            "总结三类最常见硬件链路故障。",
        ],
    },
    {
        "id": 5,
        "slug": "lecture-05-realtime-lab",
        "title": "实验二：嵌入式实时性",
        "module_type": "实验",
        "theme": "中断延迟、Jitter 与系统响应",
        "objectives": [
            "测量 GPIO、PWM、ADC、DMA 路径的时延与抖动。",
            "理解采样、处理与通信链路如何共同影响实时性。",
            "掌握用时间戳和统计量描述实时系统性能的方法。",
        ],
        "foundations": [
            "中断响应与上下文切换。",
            "DMA 与 CPU 占用的权衡。",
            "抖动、平均延迟、最坏响应时间。",
        ],
        "principles": [
            "设置基准触发信号。",
            "记录多次时间戳。",
            "计算均值、方差和极值。",
            "分析瓶颈并提出改进。",
        ],
        "metrics": [
            ("平均延迟", "典型响应速度", "影响控制闭环"),
            ("最坏时延", "极端风险", "决定安全裕度"),
            ("Jitter", "时间稳定性", "影响高频采样"),
            ("CPU 占用", "系统负荷", "影响并发任务"),
        ],
        "case_title": "案例：DMA 前后实时性对比",
        "case_points": [
            "DMA 能降低 CPU 负担，但需要处理缓冲区同步问题。",
            "日志打印过多会显著增加抖动。",
            "最坏时延常常比平均值更值得关注。",
        ],
        "practice": [
            "设计并执行一次延迟测量实验。",
            "比较中断与 DMA 两种路径的实时性。",
            "输出数据图表和改进建议。",
        ],
        "homework": [
            "提交时间戳精度报告。",
            "解释一个降低 Jitter 的具体策略。",
        ],
    },
    {
        "id": 6,
        "slug": "lecture-06-gnss-rtk",
        "title": "GNSS 与高精度定位",
        "module_type": "讲授",
        "theme": "伪距、载波相位与高精度定位流程",
        "objectives": [
            "理解 GNSS 测量观测模型及误差来源。",
            "区分单点定位、RTK 和 PPP 的基本思路与适用场景。",
            "建立 GNSS/INS 紧耦合在工程系统中的位置。",
        ],
        "foundations": [
            "NMEA、RINEX 与卫星观测量。",
            "伪距、载波相位、多路径与钟差。",
            "RTK、PPP 与紧耦合定位框架。",
        ],
        "principles": [
            "获取原始观测数据。",
            "建模主要误差项。",
            "通过差分或精密产品抑制误差。",
            "与惯导数据联合提高连续性。",
        ],
        "metrics": [
            ("水平精度", "位置误差", "决定导航效果"),
            ("初始化时间", "RTK 固定效率", "影响可用性"),
            ("可用卫星数", "几何条件", "影响解算稳定性"),
            ("失锁恢复", "复杂环境鲁棒性", "影响连续运行"),
        ],
        "case_title": "案例：校园道路 RTK 轨迹对比",
        "case_points": [
            "单点定位在遮挡场景下偏差明显增加。",
            "RTK 在基站和观测条件良好时可显著提升精度。",
            "GNSS/INS 联合可平滑短时失锁段。",
        ],
        "practice": [
            "解析 NMEA 与 RINEX 数据。",
            "计算定位误差并比较不同模式。",
            "为后续 RTK 实验准备数据处理流程。",
        ],
        "homework": [
            "绘制 GNSS 误差来源示意图。",
            "说明 RTK 与 PPP 的一个关键差异。",
        ],
    },
    {
        "id": 7,
        "slug": "lecture-07-gnss-driver-lab",
        "title": "实验三：GNSS 驱动与 RTK",
        "module_type": "实验",
        "theme": "模块驱动、数据采集与后处理对比",
        "objectives": [
            "完成 GNSS 模块驱动与数据采集链路。",
            "掌握 RTKLIB 后处理的基本流程与结果判读。",
            "建立单点定位与 RTK 结果的量化比较方法。",
        ],
        "foundations": [
            "Ublox 模块数据输出配置。",
            "RTKLIB 输入输出文件结构。",
            "轨迹误差与固定率评价方法。",
        ],
        "principles": [
            "配置并采集原始数据。",
            "转换为标准处理格式。",
            "运行 RTK 解算并输出轨迹。",
            "比较不同定位模式结果。",
        ],
        "metrics": [
            ("固定解比例", "高精度状态占比", "反映观测质量"),
            ("轨迹平滑度", "结果连续性", "反映算法稳定性"),
            ("误差均值", "总体偏差", "反映最终精度"),
            ("数据完整性", "采集时长与文件正确性", "反映实验质量"),
        ],
        "case_title": "案例：单点定位与 RTK 轨迹叠加",
        "case_points": [
            "统一坐标系和采样时间后才能公平比较。",
            "高精度结果依赖良好的基线和观测环境。",
            "实验报告应给出定量误差，不只给轨迹图。",
        ],
        "practice": [
            "完成 GNSS 数据采集与后处理。",
            "导出轨迹图、误差表和固定率统计。",
            "上传原始数据与处理脚本。",
        ],
        "homework": [
            "提交 RTK 对比实验报告。",
            "解释一次定位结果退化的可能原因。",
        ],
    },
    {
        "id": 8,
        "slug": "lecture-08-imu-modeling",
        "title": "惯性传感器原理与建模",
        "module_type": "讲授",
        "theme": "MEMS 机理、误差项与标定模型",
        "objectives": [
            "理解加速度计、陀螺仪与磁力计的基本工作机理。",
            "掌握六位置标定、比例因子误差和安装误差的建模思路。",
            "建立温漂、随机噪声与积分误差之间的联系。",
        ],
        "foundations": [
            "MEMS 结构与敏感原理。",
            "12 参数标定模型与误差传递。",
            "温度、噪声与时间积分漂移。",
        ],
        "principles": [
            "采集原始静态与动态数据。",
            "建立误差模型和约束方程。",
            "求解标定参数。",
            "验证标定前后性能变化。",
        ],
        "metrics": [
            ("零偏", "静态输出偏移", "决定积分起点"),
            ("比例因子", "量程映射偏差", "决定量测斜率"),
            ("交轴误差", "轴间耦合", "影响姿态解算"),
            ("温漂", "随温度变化的偏差", "影响长期稳定性"),
        ],
        "case_title": "案例：六位置标定到姿态误差的传递",
        "case_points": [
            "标定误差会直接进入姿态和速度积分结果。",
            "仅做静态标定并不能完全消除动态误差。",
            "温度实验有助于理解工业级与消费级差异。",
        ],
        "practice": [
            "推导六位置标定方程。",
            "构建标定数据表。",
            "为实验四准备标定脚本与记录模板。",
        ],
        "homework": [
            "整理 12 参数模型的变量含义。",
            "解释一个磁力计椭球拟合的工程意义。",
        ],
    },
    {
        "id": 9,
        "slug": "lecture-09-imu-calibration-lab",
        "title": "实验四：IMU 完整标定",
        "module_type": "实验",
        "theme": "加速度、陀螺与磁力计联合标定",
        "objectives": [
            "完成 IMU 三类传感器的标定流程。",
            "用 Allan 方差、静态误差和姿态结果评价改进效果。",
            "输出具备工程说服力的标定专报。",
        ],
        "foundations": [
            "标定数据采集规范。",
            "Allan Variance 评估流程。",
            "标定前后姿态和量测误差对比。",
        ],
        "principles": [
            "采集多姿态、多温度数据。",
            "求解标定参数。",
            "评估随机噪声与长期稳定性。",
            "整理提升数据并形成专报。",
        ],
        "metrics": [
            ("零偏改善", "静态输出修正", "反映标定有效性"),
            ("姿态误差下降", "最终应用指标", "反映工程价值"),
            ("温漂抑制", "跨温性能", "反映鲁棒性"),
            ("报告完整度", "数据与图表质量", "反映表达能力"),
        ],
        "case_title": "案例：求职级标定专报的结构",
        "case_points": [
            "图表必须显示标定前后量化提升。",
            "报告要解释方法、结果和边界条件。",
            "工程化表达比堆砌公式更重要。",
        ],
        "practice": [
            "执行标定、噪声分析和结果验证。",
            "形成标定专报初稿。",
            "上传脚本、数据和图表。",
        ],
        "homework": [
            "提交求职级标定专报。",
            "回答一项仍未完全解决的误差来源。",
        ],
    },
    {
        "id": 10,
        "slug": "lecture-10-magnetic-signal-integrity",
        "title": "磁传感器与信号完整性",
        "module_type": "讲授",
        "theme": "航向测量、前端设计与抗干扰",
        "objectives": [
            "理解磁传感器测量航向的基本原理与误差来源。",
            "掌握硬铁、软铁、布线和电源噪声对磁测量的影响。",
            "理解模拟前端和 PCB Layout 在传感系统中的作用。",
        ],
        "foundations": [
            "磁场测量与姿态参考。",
            "硬铁与软铁干扰模型。",
            "抗混叠、地线回流和布局布线基础。",
        ],
        "principles": [
            "识别干扰源与耦合路径。",
            "通过标定和布局降低误差。",
            "在多环境下复测航向精度。",
            "用数据驱动 Layout 优化。",
        ],
        "metrics": [
            ("航向误差", "角度偏差", "决定导航价值"),
            ("噪声底", "短时抖动", "影响可滤波性"),
            ("环境敏感度", "对周围金属和电流的响应", "影响鲁棒性"),
            ("布线质量", "信号完整性", "影响可重复性"),
        ],
        "case_title": "案例：PCB 布局前后磁测量对比",
        "case_points": [
            "高电流路径与磁敏元件距离过近会显著恶化结果。",
            "良好的地参考和器件摆放能降低噪声。",
            "磁力计标定无法完全代替良好硬件设计。",
        ],
        "practice": [
            "分析一块采集板的干扰路径。",
            "整理航向误差来源列表。",
            "为实验五准备环境测试方案。",
        ],
        "homework": [
            "提交 Layout 案例评述。",
            "给出一种降低磁干扰的设计策略。",
        ],
    },
    {
        "id": 11,
        "slug": "lecture-11-magnetometer-lab",
        "title": "实验五：磁力计驱动与标定",
        "module_type": "实验",
        "theme": "方位角解算与环境适应性验证",
        "objectives": [
            "完成磁力计驱动、原始数据读取与航向角解算。",
            "通过椭球拟合完成标定并比较前后精度。",
            "在不同环境中验证系统鲁棒性。",
        ],
        "foundations": [
            "磁矢量与航向角关系。",
            "硬铁软铁校正。",
            "多环境测试与误差统计。",
        ],
        "principles": [
            "采集原始磁场数据。",
            "拟合椭球并求校正参数。",
            "计算航向角并与参考对比。",
            "分析不同环境下退化原因。",
        ],
        "metrics": [
            ("校正前后误差", "角度改进幅度", "衡量标定价值"),
            ("环境稳定性", "场景变化下误差", "衡量鲁棒性"),
            ("驱动完整性", "数据读取与输出可靠性", "衡量工程质量"),
            ("报告说服力", "数据图表与解释", "衡量表达质量"),
        ],
        "case_title": "案例：教室与室外环境航向对比",
        "case_points": [
            "室内金属设施和电源设备常引入显著偏差。",
            "标定后的改善应通过多点角度测试体现。",
            "环境标签是实验数据的重要元信息。",
        ],
        "practice": [
            "完成驱动、标定和多环境测试。",
            "输出误差图、散点图与改进结论。",
            "提交驱动程序和实验报告。",
        ],
        "homework": [
            "总结一个室内磁场异常案例。",
            "补充航向角误差的定量分析。",
        ],
    },
    {
        "id": 12,
        "slug": "lecture-12-pressure-sensors",
        "title": "压力传感器原理及应用",
        "module_type": "讲授",
        "theme": "压力测量、温漂补偿与高度应用",
        "objectives": [
            "理解压阻式与电容式压力测量基本原理。",
            "掌握温度补偿、非线性补偿和高度换算思路。",
            "建立压力传感器在高度估计中的位置。",
        ],
        "foundations": [
            "压力-电信号转换机理。",
            "传感器内置补偿参数与外部标定。",
            "气压与高度关系及应用边界。",
        ],
        "principles": [
            "读取原始压力和温度数据。",
            "应用补偿模型得到物理量。",
            "转换为相对高度或海拔。",
            "结合滤波抑制动态噪声。",
        ],
        "metrics": [
            ("压力精度", "静态压力误差", "决定高度分辨率"),
            ("温漂", "温度变化下偏差", "决定跨环境性能"),
            ("分辨率", "微小变化可检测性", "决定高度细节"),
            ("响应速度", "动态变化跟踪能力", "影响控制"),
        ],
        "case_title": "案例：室内楼层高度检测",
        "case_points": [
            "相对高度估计往往比绝对海拔更有实用价值。",
            "温漂与空调气流都可能污染实验数据。",
            "与 IMU 融合可补偿动态响应不足。",
        ],
        "practice": [
            "阅读 MS 系列器件手册。",
            "整理压力-高度转换关系。",
            "为实验六准备滤波模型。",
        ],
        "homework": [
            "提交器件手册精读摘要。",
            "说明压力高度估计的一个局限。",
        ],
    },
    {
        "id": 13,
        "slug": "lecture-13-barometer-lab",
        "title": "实验六：气压计高度解算",
        "module_type": "实验",
        "theme": "滤波、建模与高度精度验证",
        "objectives": [
            "完成气压计高度解算与滤波设计。",
            "比较低通滤波和卡尔曼模型在高度估计中的表现。",
            "形成可量化的高度精度分析报告。",
        ],
        "foundations": [
            "气压到高度的数学关系。",
            "低通滤波与一维卡尔曼建模。",
            "静态与动态场景下的误差评价。",
        ],
        "principles": [
            "读取并补偿气压数据。",
            "构造相对高度序列。",
            "应用滤波模型并调参。",
            "对比不同方案的精度与响应。",
        ],
        "metrics": [
            ("静态波动", "高度稳定性", "决定静止精度"),
            ("动态响应", "变化跟踪能力", "决定运动适应性"),
            ("过冲与延迟", "滤波代价", "决定控制适用性"),
            ("最终误差", "与参考值偏差", "决定工程结论"),
        ],
        "case_title": "案例：楼梯上下行高度曲线",
        "case_points": [
            "简单低通可抑噪，但会引入明显迟滞。",
            "合理建模的卡尔曼滤波兼顾稳定性和响应速度。",
            "实验报告要同时给静态和动态曲线。",
        ],
        "practice": [
            "搭建两种滤波方案并调参。",
            "采集楼梯或升降实验数据。",
            "输出曲线、参数和结论。",
        ],
        "homework": [
            "提交滤波对比报告。",
            "解释一次高度曲线异常的可能成因。",
        ],
    },
    {
        "id": 14,
        "slug": "lecture-14-eskf",
        "title": "数据融合（一）：ESKF",
        "module_type": "讲授",
        "theme": "误差状态卡尔曼滤波与姿态表示",
        "objectives": [
            "理解 ESKF 相对 EKF 的建模优势。",
            "掌握误差状态、名义状态和观测更新的基本结构。",
            "理解 SO(3) 扰动模型在姿态估计中的必要性。",
        ],
        "foundations": [
            "KF、EKF、UKF 的角色与局限。",
            "名义状态与误差状态分离。",
            "李群、李代数与小扰动更新。",
        ],
        "principles": [
            "建立名义运动模型。",
            "传播误差协方差。",
            "利用观测进行误差修正。",
            "把误差注入名义状态。",
        ],
        "metrics": [
            ("姿态误差", "估计精度", "核心输出指标"),
            ("一致性", "协方差与真实误差匹配度", "决定可信度"),
            ("数值稳定性", "长时运行表现", "决定工程可用性"),
            ("实时性", "计算复杂度", "决定嵌入式部署"),
        ],
        "case_title": "案例：IMU 与外部观测联合姿态估计",
        "case_points": [
            "ESKF 更适合处理姿态这类非线性状态。",
            "误差注入机制有助于保持数值稳定。",
            "观测模型设计决定收敛速度与精度。",
        ],
        "practice": [
            "推导误差状态方程。",
            "写出 SO(3) 小扰动更新关系。",
            "准备后续姿态算法对比实验。",
        ],
        "homework": [
            "完成 ESKF 状态方程笔记。",
            "解释名义状态和误差状态分离的好处。",
        ],
    },
    {
        "id": 15,
        "slug": "lecture-15-fusion-comparison",
        "title": "数据融合（二）：算法对比",
        "module_type": "讲授",
        "theme": "Mahony、Madgwick 与 EKF/ESKF 比较",
        "objectives": [
            "理解常见姿态融合算法的假设、优点和不足。",
            "掌握多速率传感器融合与可观测性分析的基本问题。",
            "为嵌入式实测中的算法选型建立评价框架。",
        ],
        "foundations": [
            "互补滤波与增益设计。",
            "Mahony、Madgwick 的结构特点。",
            "EKF/ESKF 的精度与复杂度权衡。",
        ],
        "principles": [
            "明确应用场景和资源约束。",
            "选择匹配的融合结构。",
            "通过实测数据比较性能。",
            "按指标而不是按偏好选算法。",
        ],
        "metrics": [
            ("静态稳定性", "长时漂移", "决定姿态基准"),
            ("动态响应", "快速运动跟踪", "决定机动性能"),
            ("计算开销", "CPU 与内存需求", "决定部署能力"),
            ("调参难度", "工程实现成本", "决定维护难度"),
        ],
        "case_title": "案例：ESP32 上三类姿态算法实测",
        "case_points": [
            "简单算法部署快，但在复杂运动下可能退化明显。",
            "滤波增益和初始标定质量同样重要。",
            "对比实验应统一数据集和评价指标。",
        ],
        "practice": [
            "整理三类算法的输入输出关系。",
            "设计一次实测比较方案。",
            "确定最终项目的算法选型标准。",
        ],
        "homework": [
            "提交算法对比表。",
            "说明何时应优先选择较轻量算法。",
        ],
    },
    {
        "id": 16,
        "slug": "lecture-16-ai-enhanced-fusion",
        "title": "数据融合（三）：AI 增强",
        "module_type": "讲授",
        "theme": "TLIO、1D-CNN 与学习型惯性导航",
        "objectives": [
            "理解 AI 在传感器去噪、特征提取和轨迹预测中的角色。",
            "认识 TLIO、DeepVIO 等方法与传统滤波的关系。",
            "建立对学习型方法训练数据、泛化与可解释性的基本判断。",
        ],
        "foundations": [
            "时序卷积、循环网络与序列建模。",
            "传统滤波与学习型残差修正的耦合方式。",
            "数据集、标签与泛化风险。",
        ],
        "principles": [
            "先定义传统方法的瓶颈。",
            "再选择 AI 模块承担特定子任务。",
            "通过离线训练和在线验证迭代。",
            "始终保留可解释的误差分析。",
        ],
        "metrics": [
            ("轨迹误差", "定位最终表现", "最核心指标"),
            ("泛化能力", "跨场景表现", "决定可迁移性"),
            ("样本效率", "训练成本", "决定落地门槛"),
            ("推理时延", "实时部署能力", "决定工程可用性"),
        ],
        "case_title": "案例：1D-CNN 去噪与 TLIO 阅读",
        "case_points": [
            "AI 更适合做传统模型难以精确刻画的残差修正。",
            "训练集与部署场景差异会直接影响效果。",
            "阅读论文时要关注输入、输出和评价协议。",
        ],
        "practice": [
            "完成 TLIO 阅读笔记。",
            "尝试构建一维时序去噪小实验。",
            "分析 AI 增强模块适合插入的位置。",
        ],
        "homework": [
            "总结一篇 AI 感知论文的核心贡献。",
            "说明一个学习型方法的工程风险。",
        ],
    },
    {
        "id": 17,
        "slug": "lecture-17-capstone-design",
        "title": "综合设计：产品级系统",
        "module_type": "综合设计",
        "theme": "Spec、PCB 与量化改进闭环",
        "objectives": [
            "能够从需求出发定义产品级传感系统规格。",
            "理解原理图、PCB、固件和验证数据的联动关系。",
            "把综合设计做成可演示、可量化、可归档的课程成果。",
        ],
        "foundations": [
            "需求分析与关键指标定义。",
            "KiCad 设计、器件选型与打样流程。",
            "Spec 验证和版本化归档。",
        ],
        "principles": [
            "先定义产品需求和指标。",
            "再完成硬件与固件实现。",
            "通过实验闭环证明改进。",
            "整理完整数字资产用于答辩。",
        ],
        "metrics": [
            ("Spec 达成率", "目标实现程度", "最终评价核心"),
            ("原型可运行性", "系统完整性", "决定演示效果"),
            ("量化提升", "前后对比数据", "决定技术说服力"),
            ("资产完整度", "文档、代码、视频是否齐备", "决定可复用性"),
        ],
        "case_title": "案例：产品级传感节点开发路线",
        "case_points": [
            "从使用场景反推精度、带宽、功耗与通信需求。",
            "PCB 与固件开发要同步考虑测试接口。",
            "没有量化数据的优化不算真正完成。",
        ],
        "practice": [
            "提交项目 Spec 和模块划分。",
            "完成硬件设计和打样准备。",
            "规划答辩所需的数据与演示视频。",
        ],
        "homework": [
            "提交项目规格说明书。",
            "列出一项计划验证的关键改进指标。",
        ],
    },
    {
        "id": 18,
        "slug": "lecture-18-defense-wrap-up",
        "title": "答辩与总结",
        "module_type": "综合设计",
        "theme": "成果展示、复盘与课程闭环",
        "objectives": [
            "完成项目演示与答辩表达。",
            "学会用数据和版本记录证明工程改进。",
            "对本课程的指标思维、实验方法和项目经验进行复盘。",
        ],
        "foundations": [
            "技术答辩结构：问题、方法、结果、边界。",
            "项目复盘：技术、流程、协作与资产沉淀。",
            "课程知识如何迁移到后续研究与工程实践。",
        ],
        "principles": [
            "清楚说明需求与方案。",
            "用数据证明改进有效。",
            "坦诚说明局限与下一步。",
            "沉淀文档以支持复用。",
        ],
        "metrics": [
            ("答辩清晰度", "表达与结构", "影响评审理解"),
            ("数据说服力", "证据链完整性", "影响技术可信度"),
            ("成果完整度", "软硬件与文档是否齐备", "影响项目评价"),
            ("复盘质量", "对问题和改进的认识", "影响持续成长"),
        ],
        "case_title": "案例：优秀课程项目的答辩结构",
        "case_points": [
            "开场应快速说明应用场景和目标指标。",
            "中段展示关键技术与改进数据。",
            "结尾给出可执行的下一步计划。",
        ],
        "practice": [
            "整理最终答辩 PPT。",
            "检查项目仓库、视频和文档完整性。",
            "完成课程复盘与自评。",
        ],
        "homework": [
            "提交完整项目归档。",
            "撰写课程复盘，说明一项最重要的能力提升。",
        ],
    },
]


def rgb_tuple(color: RGBColor) -> tuple[int, int, int]:
    return color[0], color[1], color[2]


def make_banner(lecture_id: int, slug: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / f"{slug}.png"
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), rgb_tuple(COLORS["sand"]))
    draw = ImageDraw.Draw(image)

    for index in range(height):
        ratio = index / height
        r = int(18 + (36 - 18) * ratio)
        g = int(60 + (105 - 60) * ratio)
        b = int(66 + (116 - 66) * ratio)
        draw.line((0, index, width, index), fill=(r, g, b))

    accent = rgb_tuple(COLORS["gold"])
    pale = (255, 248, 232)
    teal = rgb_tuple(COLORS["teal"])

    draw.ellipse((80, 100, 520, 540), outline=pale, width=12)
    draw.ellipse((980, 120, 1460, 600), outline=accent, width=14)
    draw.rounded_rectangle((180, 540, 520, 760), radius=40, outline=pale, width=8)
    draw.rounded_rectangle((840, 590, 1250, 800), radius=40, outline=(184, 225, 226), width=8)

    points = [
        (240, 220),
        (410, 320),
        (620, 250),
        (780, 430),
        (980, 350),
        (1190, 470),
        (1380, 390),
    ]
    for start, end in zip(points, points[1:]):
        draw.line((start, end), fill=accent, width=8)
    for point in points:
        draw.ellipse((point[0] - 16, point[1] - 16, point[0] + 16, point[1] + 16), fill=pale)
        draw.ellipse((point[0] - 10, point[1] - 10, point[0] + 10, point[1] + 10), fill=accent)

    base_x = 180 + lecture_id * 25
    for idx in range(6):
        left = base_x + idx * 42
        draw.rectangle((left, 680 - idx * 25, left + 24, 760), fill=(255, 255, 255, 70), outline=None)

    draw.arc((1120, 560, 1480, 920), start=205, end=336, fill=accent, width=10)
    draw.line((1220, 810, 1340, 690), fill=pale, width=10)
    draw.line((1340, 690, 1410, 760), fill=pale, width=10)

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)
    return path


def set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=22, bold=False,
                color: RGBColor = COLORS["ink"], align=PP_ALIGN.LEFT,
                font_name="Microsoft YaHei", margin=0.08, level=0):
    box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = box.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(margin)
    text_frame.margin_right = Inches(margin)
    text_frame.margin_top = Inches(margin)
    text_frame.margin_bottom = Inches(margin)
    p = text_frame.paragraphs[0]
    p.text = text
    p.level = level
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    return box


def add_bullets(slide, left, top, width, height, bullets: Iterable[str], font_size=20,
                color: RGBColor = COLORS["ink"], bullet_color: RGBColor = COLORS["gold"]):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.bullet = True
        p.level = 0
        p.space_after = Pt(8)
        run = p.runs[0]
        run.font.size = Pt(font_size)
        run.font.name = "Microsoft YaHei"
        run.font.color.rgb = color
    return box


def add_title_band(slide, lecture) -> None:
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.72))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["teal"]
    bar.line.fill.background()
    add_textbox(slide, Inches(0.45), Inches(0.12), Inches(8), Inches(0.4), lecture["title"], font_size=28, bold=True, color=COLORS["white"])
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(10.55), Inches(0.12), Inches(2.2), Inches(0.4))
    tag.fill.solid()
    tag.fill.fore_color.rgb = COLORS["gold"]
    tag.line.fill.background()
    add_textbox(slide, Inches(10.72), Inches(0.16), Inches(1.85), Inches(0.25), lecture["module_type"], font_size=18, bold=True, color=COLORS["teal_dark"], align=PP_ALIGN.CENTER)


def add_footer(slide, lecture) -> None:
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.42), Inches(7.0), Inches(12.9), Inches(7.0))
    line.line.color.rgb = COLORS["line"]
    line.line.width = Pt(1.2)
    add_textbox(slide, Inches(0.45), Inches(7.02), Inches(5), Inches(0.25), f"{COURSE_TITLE} | {COURSE_SUBTITLE}", font_size=10, color=COLORS["muted"])
    add_textbox(slide, Inches(11.2), Inches(7.02), Inches(1.4), Inches(0.25), f"L{lecture['id']:02d}", font_size=10, color=COLORS["muted"], align=PP_ALIGN.RIGHT)


def add_card(slide, left, top, width, height, title, lines, fill_color=COLORS["white"]) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = COLORS["line"]
    add_textbox(slide, left + Inches(0.12), top + Inches(0.08), width - Inches(0.24), Inches(0.38), title, font_size=18, bold=True, color=COLORS["teal"])
    add_bullets(slide, left + Inches(0.08), top + Inches(0.48), width - Inches(0.16), height - Inches(0.56), lines, font_size=16)


def slide_title(prs, lecture, banner_path: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    ribbon = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = COLORS["sand"]
    ribbon.line.fill.background()
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(4.2), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLORS["teal"]
    accent.line.fill.background()
    add_textbox(slide, Inches(0.55), Inches(0.6), Inches(3), Inches(0.35), COURSE_TITLE, font_size=20, bold=True, color=COLORS["gold"])
    add_textbox(slide, Inches(0.55), Inches(1.2), Inches(3.15), Inches(1.4), lecture["title"], font_size=30, bold=True, color=COLORS["white"])
    add_textbox(slide, Inches(0.55), Inches(2.8), Inches(3.05), Inches(1.3), lecture["theme"], font_size=18, color=COLORS["white"])
    chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(4.45), Inches(1.75), Inches(0.42))
    chip.fill.solid()
    chip.fill.fore_color.rgb = COLORS["gold"]
    chip.line.fill.background()
    add_textbox(slide, Inches(0.68), Inches(4.53), Inches(1.45), Inches(0.2), lecture["module_type"], font_size=16, bold=True, color=COLORS["teal_dark"], align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0.55), Inches(5.15), Inches(2.6), Inches(0.5), f"第 {lecture['id']:02d} 讲", font_size=16, color=COLORS["mist"])
    slide.shapes.add_picture(str(banner_path), Inches(3.95), Inches(0.55), Inches(8.85), Inches(6.1))
    add_textbox(slide, Inches(4.35), Inches(6.78), Inches(7.9), Inches(0.3), "官方授课计划驱动 | 图文并茂版课程课件", font_size=12, color=COLORS["muted"])
    add_footer(slide, lecture)


def slide_objectives(prs, lecture):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    add_title_band(slide, lecture)
    add_textbox(slide, Inches(0.55), Inches(0.95), Inches(5.4), Inches(0.4), "学习目标", font_size=24, bold=True, color=COLORS["teal"])
    add_bullets(slide, Inches(0.6), Inches(1.42), Inches(5.8), Inches(3.85), lecture["objectives"], font_size=20)

    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.15), Inches(5.45), Inches(4.6))
    panel.fill.solid()
    panel.fill.fore_color.rgb = COLORS["white"]
    panel.line.color.rgb = COLORS["line"]
    add_textbox(slide, Inches(7.45), Inches(1.42), Inches(4.9), Inches(0.35), "能力地图", font_size=22, bold=True, color=COLORS["teal"])

    cards = [
        ("理解", lecture["objectives"][0]),
        ("分析", lecture["objectives"][1]),
        ("实施", lecture["objectives"][2]),
    ]
    top_positions = [2.1, 3.1, 4.1]
    for (title, text), top in zip(cards, top_positions):
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.45), Inches(top), Inches(4.55), Inches(0.72))
        box.fill.solid()
        box.fill.fore_color.rgb = COLORS["mist"]
        box.line.color.rgb = COLORS["line"]
        add_textbox(slide, Inches(7.64), Inches(top + 0.08), Inches(1), Inches(0.24), title, font_size=15, bold=True, color=COLORS["gold"])
        add_textbox(slide, Inches(8.55), Inches(top + 0.06), Inches(3.15), Inches(0.42), text, font_size=13, color=COLORS["ink"])

    add_footer(slide, lecture)


def slide_foundation(prs, lecture):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    add_title_band(slide, lecture)
    add_textbox(slide, Inches(0.55), Inches(0.95), Inches(6), Inches(0.4), "理论基础", font_size=24, bold=True, color=COLORS["teal"])
    positions = [0.6, 4.35, 8.1]
    titles = ["概念框架", "关键机理", "工程约束"]
    for left, title, content in zip(positions, titles, lecture["foundations"]):
        add_card(slide, Inches(left), Inches(1.55), Inches(3.15), Inches(4.6), title, [content], fill_color=COLORS["white"])
    add_footer(slide, lecture)


def slide_principles(prs, lecture):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    add_title_band(slide, lecture)
    add_textbox(slide, Inches(0.55), Inches(0.95), Inches(6), Inches(0.4), "工作原理 / 实施流程", font_size=24, bold=True, color=COLORS["teal"])

    step_left = [0.7, 3.45, 6.2, 8.95]
    for idx, (left, step) in enumerate(zip(step_left, lecture["principles"]), start=1):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(left), Inches(2.2), Inches(2.1), Inches(1.15))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["gold"] if idx % 2 else COLORS["teal"]
        shape.line.fill.background()
        add_textbox(slide, Inches(left + 0.15), Inches(2.42), Inches(1.8), Inches(0.5), f"{idx}. {step}", font_size=16, bold=True, color=COLORS["white"])

    note = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.8), Inches(1.55))
    note.fill.solid()
    note.fill.fore_color.rgb = COLORS["white"]
    note.line.color.rgb = COLORS["line"]
    add_textbox(slide, Inches(1.0), Inches(4.45), Inches(11.2), Inches(0.9), "课堂提示：流程图用于帮助学生把“原理、实现、验证”连成闭环，授课时应强调每一步对应的数据证据和常见失败点。", font_size=18, color=COLORS["ink"])
    add_footer(slide, lecture)


def slide_metrics(prs, lecture):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    add_title_band(slide, lecture)
    add_textbox(slide, Inches(0.55), Inches(0.95), Inches(6), Inches(0.4), "关键指标与 Datasheet 关注点", font_size=24, bold=True, color=COLORS["teal"])
    table = slide.shapes.add_table(5, 3, Inches(0.75), Inches(1.6), Inches(7.6), Inches(3.9)).table
    headers = ["指标", "解释", "工程意义"]
    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS["teal"]
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.name = "Microsoft YaHei"
                run.font.color.rgb = COLORS["white"]
                run.font.size = Pt(16)
    for row, metric in enumerate(lecture["metrics"], start=1):
        for col, text in enumerate(metric):
            cell = table.cell(row, col)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS["white"] if row % 2 else COLORS["mist"]
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(14)
                    run.font.color.rgb = COLORS["ink"]
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(1.6), Inches(3.75), Inches(3.9))
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS["white"]
    card.line.color.rgb = COLORS["line"]
    add_textbox(slide, Inches(8.95), Inches(1.85), Inches(3.15), Inches(0.4), "教师提醒", font_size=20, bold=True, color=COLORS["gold"])
    add_bullets(slide, Inches(8.95), Inches(2.35), Inches(3.1), Inches(2.7), [
        "不要只解释参数含义，要解释它如何影响系统级结果。",
        "鼓励学生把手册参数与后续实验的测量结果对应起来。",
        "强调典型值、最大值和测试条件的区别。",
    ], font_size=16)
    add_footer(slide, lecture)


def slide_case(prs, lecture):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    add_title_band(slide, lecture)
    add_textbox(slide, Inches(0.55), Inches(0.95), Inches(6), Inches(0.4), "工程案例", font_size=24, bold=True, color=COLORS["teal"])
    hero = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.6), Inches(5.2), Inches(4.8))
    hero.fill.solid()
    hero.fill.fore_color.rgb = COLORS["teal"]
    hero.line.fill.background()
    add_textbox(slide, Inches(1.05), Inches(1.95), Inches(4.4), Inches(0.5), lecture["case_title"], font_size=24, bold=True, color=COLORS["gold"])
    add_bullets(slide, Inches(1.02), Inches(2.65), Inches(4.45), Inches(2.85), lecture["case_points"], font_size=18, color=COLORS["white"])

    add_card(slide, Inches(6.35), Inches(1.6), Inches(2.75), Inches(2.1), "输入", [lecture["foundations"][0], lecture["foundations"][1]], fill_color=COLORS["white"])
    add_card(slide, Inches(9.4), Inches(1.6), Inches(2.75), Inches(2.1), "输出", [lecture["practice"][0], lecture["practice"][1]], fill_color=COLORS["white"])
    add_card(slide, Inches(6.35), Inches(4.0), Inches(5.8), Inches(2.4), "可视化讲授建议", [
        "在这一页配合板书或实物，帮助学生把抽象算法映射到真实系统。",
        "尽量给出一个量化结果，如误差下降、轨迹变稳、响应加快。",
    ], fill_color=COLORS["mist"])
    add_footer(slide, lecture)


def slide_practice(prs, lecture):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    add_title_band(slide, lecture)
    add_textbox(slide, Inches(0.55), Inches(0.95), Inches(6), Inches(0.4), "实践路径与课堂活动", font_size=24, bold=True, color=COLORS["teal"])
    phases = ["准备", "实施", "沉淀"]
    lefts = [0.8, 4.35, 7.9]
    colors = [COLORS["gold"], COLORS["teal"], COLORS["gold"]]
    for idx, (phase, left, fill_color, content) in enumerate(zip(phases, lefts, colors, lecture["practice"])):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.8), Inches(2.95), Inches(4.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.fill.background()
        add_textbox(slide, Inches(left + 0.2), Inches(2.05), Inches(2.2), Inches(0.3), phase, font_size=22, bold=True, color=COLORS["white"])
        add_textbox(slide, Inches(left + 0.2), Inches(2.7), Inches(2.35), Inches(2.7), content, font_size=18, color=COLORS["white"])
        index_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(left + 2.2), Inches(1.95), Inches(0.5), Inches(0.5))
        index_box.fill.solid()
        index_box.fill.fore_color.rgb = COLORS["white"]
        index_box.line.fill.background()
        add_textbox(slide, Inches(left + 2.31), Inches(2.04), Inches(0.26), Inches(0.18), str(idx + 1), font_size=16, bold=True, color=fill_color, align=PP_ALIGN.CENTER)
    add_footer(slide, lecture)


def slide_summary(prs, lecture):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, COLORS["sand"])
    add_title_band(slide, lecture)
    add_textbox(slide, Inches(0.55), Inches(0.95), Inches(5), Inches(0.4), "总结与作业", font_size=24, bold=True, color=COLORS["teal"])
    add_card(slide, Inches(0.75), Inches(1.6), Inches(5.55), Inches(4.9), "本讲关键结论", [
        lecture["objectives"][0],
        lecture["objectives"][1],
        lecture["objectives"][2],
    ], fill_color=COLORS["white"])
    add_card(slide, Inches(6.7), Inches(1.6), Inches(5.55), Inches(4.9), "课后任务", lecture["homework"], fill_color=COLORS["mist"])
    add_footer(slide, lecture)


def generate_notes(lecture) -> None:
    NOTES_DIR.mkdir(parents=True, exist_ok=True)
    content = [
        f"# {lecture['title']}",
        "",
        f"- 讲次编号：{lecture['id']:02d}",
        f"- 类型：{lecture['module_type']}",
        f"- 主题：{lecture['theme']}",
        "",
        "## Slide 1 标题页讲稿",
        f"强调本讲在整门课程中的位置，突出“{lecture['theme']}”与官方授课计划的对应关系。",
        "",
        "## Slide 2 学习目标讲稿",
        *[f"- {item}" for item in lecture["objectives"]],
        "",
        "## Slide 3 理论基础讲稿",
        *[f"- {item}" for item in lecture["foundations"]],
        "",
        "## Slide 4 工作原理讲稿",
        *[f"- 步骤 {idx}: {item}" for idx, item in enumerate(lecture["principles"], start=1)],
        "",
        "## Slide 5 指标解读讲稿",
        *[f"- {name}：{meaning}，工程意义是{impact}。" for name, meaning, impact in lecture["metrics"]],
        "",
        "## Slide 6 案例讲稿",
        f"- 案例主题：{lecture['case_title']}",
        *[f"- {item}" for item in lecture["case_points"]],
        "",
        "## Slide 7 实践活动讲稿",
        *[f"- {item}" for item in lecture["practice"]],
        "",
        "## Slide 8 总结与作业讲稿",
        *[f"- {item}" for item in lecture["homework"]],
        "",
        "## 教学提示",
        "- 先讲指标与目标，再讲实现与验证，避免只讲概念不讲工程闭环。",
        "- 尽量在课堂上展示一张数据图、一段代码接口或一张系统框图。",
    ]
    (NOTES_DIR / f"{lecture['slug']}.md").write_text("\n".join(content), encoding="utf-8")


def generate_manifest() -> None:
    lines = [
        "# 2026 Spring PPT Index",
        "",
        "本目录保存整学期生成的 PPT 与配套讲稿。",
        "",
        "## 目录",
        "",
        "- pptx/：18 个课程单元的 PowerPoint 文件",
        "- notes/：对应讲稿与讲者备注",
        "- assets/：课件内使用的抽象插图资源",
        "",
        "## 清单",
        "",
    ]
    for lecture in LECTURES:
        lines.append(f"- {lecture['slug']}.pptx：{lecture['title']}")
    (ROOT / "lectures" / "2026-spring" / "PPT_INDEX.md").write_text("\n".join(lines), encoding="utf-8")


def generate_presentation(lecture) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    banner = make_banner(lecture["id"], lecture["slug"])
    slide_title(prs, lecture, banner)
    slide_objectives(prs, lecture)
    slide_foundation(prs, lecture)
    slide_principles(prs, lecture)
    slide_metrics(prs, lecture)
    slide_case(prs, lecture)
    slide_practice(prs, lecture)
    slide_summary(prs, lecture)

    prs.save(OUTPUT_DIR / f"{lecture['slug']}.pptx")
    generate_notes(lecture)


def main() -> None:
    for lecture in LECTURES:
        generate_presentation(lecture)
    generate_manifest()
    print(f"Generated {len(LECTURES)} presentations in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()